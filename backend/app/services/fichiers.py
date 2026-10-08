"""Génération de fichiers tableur (XLSX, CSV) pour virements, déclarations et comptabilité."""
import csv
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO, StringIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _valeur(v):
    if isinstance(v, Decimal):
        return float(v)
    return v


def xlsx(titre: str, entetes: list, lignes: list, totaux: dict | None = None, sous_titre: str | None = None) -> bytes:
    """`totaux` : {nom de colonne: valeur} -> une ligne de total en gras sous le tableau."""
    wb = Workbook()
    ws = wb.active
    ws.title = titre[:31]
    ligne0 = 1
    if sous_titre:
        ws.cell(row=1, column=1, value=sous_titre).font = Font(bold=True, size=12)
        ligne0 = 3
    for j, e in enumerate(entetes, 1):
        c = ws.cell(row=ligne0, column=j, value=e)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, l in enumerate(lignes, ligne0 + 1):
        for j, v in enumerate(l, 1):
            c = ws.cell(row=i, column=j, value=_valeur(v))
            if isinstance(c.value, float):
                c.number_format = "#,##0.00" if c.value != int(c.value) else "#,##0"
            if isinstance(v, (date, datetime)):
                c.number_format = "DD/MM/YYYY"
    if totaux:
        r = ligne0 + len(lignes) + 1
        ws.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
        for nom, v in totaux.items():
            c = ws.cell(row=r, column=entetes.index(nom) + 1, value=_valeur(v))
            c.font, c.number_format = Font(bold=True), "#,##0.00"
    for j, e in enumerate(entetes, 1):
        largeur = max([len(str(e))] + [len(str(_valeur(l[j - 1]) if j - 1 < len(l) else "")) for l in lignes[:200]]) + 2
        ws.column_dimensions[get_column_letter(j)].width = min(max(largeur, 10), 45)
    ws.freeze_panes = ws.cell(row=ligne0 + 1, column=1)
    tampon = BytesIO()
    wb.save(tampon)
    return tampon.getvalue()


def csv_fr(entetes: list, lignes: list) -> bytes:
    """CSV lisible directement par Excel en français : séparateur « ; », virgule décimale, BOM UTF-8."""
    s = StringIO()
    w = csv.writer(s, delimiter=";", lineterminator="\r\n")
    w.writerow(entetes)
    for l in lignes:
        w.writerow([str(v).replace(".", ",") if isinstance(v, (Decimal, float)) else ("" if v is None else v) for v in l])
    return ("﻿" + s.getvalue()).encode("utf-8")


MEDIA = {"xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "csv": "text/csv; charset=utf-8"}
